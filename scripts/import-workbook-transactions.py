#!/usr/bin/env python3
"""Import the workbook's Transactions sheet into a copied Pocket Guard SQLite database."""

import argparse
import hashlib
import re
import sqlite3
import sys
import uuid
from datetime import date, datetime, time, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

DATE_HEADERS = {"date", "transactiondate", "occurredon", "occurredat"}
CATEGORY_HEADERS = {"category", "budgetline", "budgetcategory", "line"}
AMOUNT_HEADERS = {"amount", "amountxaf", "value", "transactionamount"}
NOTE_HEADERS = {"note", "description", "details", "particulars", "narration"}
DOUALA = ZoneInfo("Africa/Douala")


def normalize_header(value):
    return re.sub(r"[^a-z0-9]", "", str(value or "").strip().lower())


def find_header(rows):
    for index, row in enumerate(rows[:50]):
        headers = {normalize_header(value): offset for offset, value in enumerate(row)}
        date_column = next((headers[key] for key in DATE_HEADERS if key in headers), None)
        category_column = next(
            (headers[key] for key in CATEGORY_HEADERS if key in headers), None
        )
        amount_column = next((headers[key] for key in AMOUNT_HEADERS if key in headers), None)
        if date_column is not None and category_column is not None and amount_column is not None:
            note_column = next((headers[key] for key in NOTE_HEADERS if key in headers), None)
            return index, date_column, category_column, amount_column, note_column
    raise ValueError("Could not find a header row with Date, Category, and Amount columns.")


def parse_date(value, row_number):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    if not text:
        raise ValueError(f"Row {row_number}: date is missing.")
    for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(text[:10], pattern).date()
        except ValueError:
            pass
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).date()
    except ValueError as error:
        raise ValueError(f"Row {row_number}: invalid date {text!r}.") from error


def parse_amount(value, row_number):
    if isinstance(value, bool) or value is None:
        raise ValueError(f"Row {row_number}: amount is invalid.")
    text = str(value).strip().replace(",", "").replace(" ", "").replace("XAF", "")
    try:
        amount = Decimal(text)
    except InvalidOperation as error:
        raise ValueError(f"Row {row_number}: invalid amount {value!r}.") from error
    if not amount.is_finite() or amount <= 0 or amount != amount.to_integral_value():
        raise ValueError(f"Row {row_number}: amount must be a positive whole XAF amount.")
    return int(amount)


def import_rows(database, rows, sheet_name="Transactions"):
    header_index, date_column, category_column, amount_column, note_column = find_header(rows)
    household_rows = database.execute(
        "SELECT id FROM households WHERE deleted_at IS NULL ORDER BY created_at LIMIT 2"
    ).fetchall()
    if len(household_rows) != 1:
        raise ValueError("The database must contain exactly one active household.")
    household_id = household_rows[0][0]

    line_rows = database.execute(
        "SELECT id, name FROM budget_lines WHERE household_id = ? "
        "AND deleted_at IS NULL AND archived = 0",
        (household_id,),
    ).fetchall()
    lines = {}
    for line_id, name in line_rows:
        key = " ".join(name.casefold().split())
        if key in lines:
            raise ValueError(f"Budget line name {name!r} is ambiguous; resolve it before importing.")
        lines[key] = line_id

    timestamp = datetime.now(timezone.utc).isoformat(timespec="milliseconds").replace("+00:00", "Z")
    inserted = 0
    skipped = 0
    duplicates = 0
    for offset, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        category = row[category_column] if category_column < len(row) else None
        amount_value = row[amount_column] if amount_column < len(row) else None
        if category is None or not str(category).strip() or amount_value is None or not str(amount_value).strip():
            skipped += 1
            continue

        key = " ".join(str(category).casefold().split())
        line_id = lines.get(key)
        if line_id is None:
            raise ValueError(f"Row {offset}: budget line {category!r} does not exist in this household.")
        date_value = row[date_column] if date_column < len(row) else None
        occurred_date = parse_date(date_value, offset)
        amount = parse_amount(amount_value, offset)
        note_value = row[note_column] if note_column is not None and note_column < len(row) else None
        note = str(note_value).strip() if note_value is not None else ""
        occurred_at = datetime.combine(occurred_date, time(12), DOUALA).astimezone(
            timezone.utc
        ).isoformat(timespec="milliseconds").replace("+00:00", "Z")
        stable_source = f"{household_id}\0{sheet_name}\0{offset}".encode()
        transaction_id = str(uuid.UUID(bytes=hashlib.sha256(stable_source).digest()[:16], version=4))
        outbox_id = str(uuid.uuid5(uuid.NAMESPACE_URL, f"pocket-guard-import:{transaction_id}"))

        cursor = database.execute(
            """INSERT OR IGNORE INTO transactions
              (id, household_id, occurred_at, amount, direction, line_id, debt_id, note,
               source, raw_text, status, override_reason, device_id, created_at, updated_at, deleted_at)
              VALUES (?, ?, ?, ?, 'expense', ?, NULL, ?, 'manual', NULL, 'confirmed', NULL, NULL, ?, ?, NULL)""",
            (transaction_id, household_id, occurred_at, amount, line_id, note or None, timestamp, timestamp),
        )
        if cursor.rowcount == 0:
            duplicates += 1
            continue
        database.execute(
            """INSERT INTO outbox
              (id, household_id, table_name, row_id, op, device_id, created_at, updated_at, deleted_at)
              VALUES (?, ?, 'transactions', ?, 'create', NULL, ?, ?, NULL)""",
            (outbox_id, household_id, transaction_id, timestamp, timestamp),
        )
        inserted += 1
    return inserted, skipped, duplicates


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path, help="Path to the owner's .xlsx workbook")
    parser.add_argument("database", type=Path, help="Path to a copied Pocket Guard SQLite database")
    parser.add_argument("--sheet", default="Transactions", help="Worksheet name (default: Transactions)")
    args = parser.parse_args()
    if not args.workbook.is_file():
        parser.error(f"Workbook not found: {args.workbook}")
    if not args.database.is_file():
        parser.error(f"Database not found: {args.database}")

    workbook = load_workbook(args.workbook, data_only=True, read_only=True)
    if args.sheet not in workbook.sheetnames:
        parser.error(f"Worksheet {args.sheet!r} not found in workbook.")
    rows = list(workbook[args.sheet].iter_rows(values_only=True))

    try:
        with sqlite3.connect(args.database) as database:
            inserted, skipped, duplicates = import_rows(database, rows, args.sheet)
    except (sqlite3.Error, ValueError) as error:
        print(f"Import failed: {error}", file=sys.stderr)
        return 1

    print(
        f"Imported {inserted} transaction(s); skipped {skipped} incomplete row(s); "
        f"already imported {duplicates} row(s)."
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
